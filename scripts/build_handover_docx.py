# -*- coding: utf-8 -*-
"""Generate project handover Word document for image_db."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
# ASCII filename avoids Windows console/codepage issues; Chinese title is inside the doc.
OUT = ROOT / "docs" / "project-handover.docx"


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, name="微软雅黑")
    return h


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    set_run_font(run, size=9, name="Consolas")
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=9)
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("图像路径式数据库管理系统")
    set_run_font(r, size=22, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("项目交接文档")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "仓库：https://github.com/luoshirui-192/image.git\n"
        "文档版本：2026-07-29\n"
        "适用分支：main / deploy/machine-a / deploy/machine-b"
    )
    set_run_font(r, size=11)
    doc.add_page_break()

    # TOC-like outline
    add_heading(doc, "目录", 1)
    add_numbered(
        doc,
        [
            "项目概述与目标",
            "技术栈与系统架构",
            "目录结构说明",
            "Git 分支与部署模式",
            "环境变量与配置",
            "部署与运维操作",
            "本地开发环境",
            "功能模块与关键文件",
            "如何修改：常见开发场景",
            "后台任务、调度与故障处理",
            "测试与验收",
            "安全与注意事项",
            "交接检查清单",
        ],
    )
    doc.add_page_break()

    # 1
    add_heading(doc, "1. 项目概述与目标", 1)
    add_para(
        doc,
        "本系统是「图像路径式数据库管理系统」：数据库中保存图片相对路径与元数据，"
        "真实文件存放在本地 upload/ 目录或 MinIO 对象存储中，不把图片以 BLOB 长期堆在业务库里。"
        "面向旧库 BLOB 迁移、模拟数据库浏览、指纹对比与业务评测等场景。",
    )
    add_heading(doc, "1.1 核心能力", 2)
    add_table(
        doc,
        ["功能", "说明", "权限"],
        [
            ["图片上传", "拖拽/多选上传，分类与标签", "登录用户"],
            ["指纹对比", "业务表样本浏览、ZIP 导入、细节点叠加", "登录用户"],
            ["指纹评测指标", "DET/ROC 等（菜单隐藏，可直达 /fingerprint-eval）", "登录用户"],
            ["BLOB 迁移", "旧库 BLOB → upload 路径表", "登录用户"],
            ["模拟数据库", "目录树、建配置、SQL、多 BLOB 预览、路径导出", "登录用户"],
            ["任务台", "迁移 / 路径导出 / 指纹导入任务监控", "登录用户"],
            ["分类管理", "维护图片分类", "登录用户"],
            ["操作日志", "SQL 与上传/删除审计", "仅管理员"],
            ["系统设置", "上传限制、SQL 超时等", "仅管理员"],
        ],
    )
    add_para(doc, "默认管理员账号：admin / admin123（交接后务必立刻修改密码）。", bold=True)

    # 2
    add_heading(doc, "2. 技术栈与系统架构", 1)
    add_heading(doc, "2.1 技术栈", 2)
    add_table(
        doc,
        ["层级", "技术"],
        [
            ["前端", "Vue 3 + Vite + Element Plus + Pinia + Vue Router + Axios"],
            ["后端", "Django 5 + DRF + SimpleJWT + Gunicorn"],
            ["数据库", "MySQL 8.0（Docker）；兼容 MySQL 5.1 的自定义 backend"],
            ["存储", "本地 upload/ 或 MinIO（S3 API）"],
            ["网关", "Nginx（托管前端 dist + 反代 /api）"],
            ["容器", "Docker Compose：db / backend / web / scheduler"],
        ],
    )
    add_heading(doc, "2.2 请求链路", 2)
    add_code(
        doc,
        "浏览器 → Nginx(web) → Gunicorn(backend/Django) → MySQL\n"
        "                 ↘ scheduler（后台任务与定时维护）\n"
        "图片访问统一走 /api/images/*（JWT 或短期 token），不直接暴露 upload/",
    )
    add_heading(doc, "2.3 双机生产拓扑（当前常见）", 2)
    add_code(
        doc,
        "浏览器 → 机器 A（应用 + MySQL + scheduler）\n"
        "              └── S3 API → MinIO VIP 192.168.9.9:9000\n"
        "                            bucket=biox  prefix=data/image_db\n"
        "机器 B：仅负责 MinIO 前缀初始化脚本（不跑业务应用）",
    )

    # 3
    add_heading(doc, "3. 目录结构说明", 1)
    add_table(
        doc,
        ["路径", "说明"],
        [
            ["backend/", "Django 后端（users / images / fingerprints / sqlquery / logs）"],
            ["backend/config/", "settings、urls、健康检查、MySQL 兼容 backend"],
            ["backend/utils/", "存储抽象、路径规范、文件安全"],
            ["frontend/src/views/", "页面组件"],
            ["frontend/src/api/", "前端 API 封装"],
            ["frontend/src/config/menu.js", "侧边栏菜单（改菜单先改这里）"],
            ["frontend/src/router/", "路由与登录/管理员守卫"],
            ["docker/", "Dockerfile、entrypoint、nginx、maintenance-loop"],
            ["deploy/", "裸机 Nginx/Gunicorn/备份模板"],
            ["scripts/", "运维、备份、冒烟测试、密钥生成"],
            ["sql/", "建库/种子/BLOB 迁移参考 DDL"],
            ["docs/", "补充文档（含本文档）"],
            ["upload/", "本地存储模式下的图片根目录（运行时）"],
            ["templates/", "指纹 ISO 模板等"],
            ["docker-compose.yml", "单机一键部署（main）"],
            ["docker-compose.app.yml", "机器 A 应用部署"],
            ["start.sh / start.ps1", "单机启动"],
            ["start-app.sh", "机器 A 启动"],
            ["start-storage.sh", "机器 B MinIO 初始化"],
        ],
    )

    # 4
    add_heading(doc, "4. Git 分支与部署模式", 1)
    add_table(
        doc,
        ["分支", "用途", "启动方式", "环境模板"],
        [
            ["main", "单机全量（本地 upload）", "start.sh / start.ps1 + docker-compose.yml", ".env.docker.example"],
            ["deploy/machine-a", "生产应用机 A（可接 MinIO）", "start-app.sh + docker-compose.app.yml", ".env.app.example"],
            ["deploy/machine-b", "仅存储初始化（不跑 App）", "start-storage.sh", ".env.storage.example"],
        ],
    )
    add_para(doc, "日常开发与功能迭代以 main 或 deploy/machine-a 为准；machine-b 代码面较窄，主要是 MinIO 初始化说明与脚本。")
    add_heading(doc, "4.1 克隆示例", 2)
    add_code(
        doc,
        "# 单机\ngit clone https://github.com/luoshirui-192/image.git\n\n"
        "# 机器 A\ngit clone -b deploy/machine-a https://github.com/luoshirui-192/image.git\n\n"
        "# 机器 B\ngit clone -b deploy/machine-b https://github.com/luoshirui-192/image.git",
    )

    # 5
    add_heading(doc, "5. 环境变量与配置", 1)
    add_para(doc, "所有密钥与环境相关配置都在 .env（勿提交 Git）。后端统一从 backend/config/settings.py 读取。")
    add_heading(doc, "5.1 配置文件位置", 2)
    add_table(
        doc,
        ["文件", "用途"],
        [
            [".env（根目录）", "Docker Compose 使用，gitignore"],
            [".env.docker.example", "单机模板"],
            [".env.app.example", "机器 A 模板"],
            [".env.storage.example", "机器 B 模板"],
            ["backend/.env.example", "本地开发后端"],
            ["frontend/.env.example", "本地开发前端（VITE_API_BASE_URL=/api）"],
            ["deploy/paths.env.example", "裸机路径渲染"],
        ],
    )
    add_heading(doc, "5.2 必改/关键变量", 2)
    add_table(
        doc,
        ["变量", "说明"],
        [
            ["PUBLIC_URL", "浏览器访问根地址，无末尾斜杠；影响 ALLOWED_HOSTS/CORS"],
            ["SECRET_KEY", "Django 密钥"],
            ["IMAGE_ACCESS_SECRET", "图片访问 token 签名密钥"],
            ["MYSQL_* / DB_*", "数据库连接"],
            ["DB_HOST", "compose 内用 db；外挂 MySQL 常用 host.docker.internal"],
            ["STORAGE_BACKEND", "local 或 minio"],
            ["MINIO_ENDPOINT", "如 http://192.168.9.9:9000"],
            ["MINIO_ACCESS_KEY / SECRET_KEY", "MinIO 凭据"],
            ["MINIO_BUCKET / MINIO_PREFIX", "通常 biox / data/image_db"],
            ["HTTP_PORT", "对外端口，默认 80"],
            ["GUNICORN_WORKERS / TIMEOUT", "并发与超时调优"],
            ["LOG_RETENTION_DAYS", "操作日志保留"],
            ["DELETED_IMAGE_RETENTION_DAYS", "软删文件清理"],
            ["MAINTENANCE_INTERVAL_HOURS", "定时维护间隔"],
        ],
    )
    add_para(doc, "生成密钥：")
    add_code(doc, "python scripts/generate_secret_key.py")
    add_para(
        doc,
        "注意：PUBLIC_URL 必须与用户浏览器地址一致，否则登录可能 403（CORS/Hosts）。"
        "修改后需重启/重新执行启动脚本，让 docker/set-env.py 同步 ALLOWED_HOSTS。",
        bold=True,
    )

    # 6
    add_heading(doc, "6. 部署与运维操作", 1)
    add_heading(doc, "6.1 单机 Docker（main）", 2)
    add_code(
        doc,
        "cp .env.docker.example .env\n"
        "# 编辑 PUBLIC_URL、密钥、MYSQL_*\n"
        "./start.sh          # Linux\n"
        ".\\start.ps1         # Windows PowerShell\n"
        "浏览器打开 PUBLIC_URL，默认 admin / admin123",
    )
    add_heading(doc, "6.2 机器 A（应用）", 2)
    add_code(
        doc,
        "cp .env.app.example .env\n"
        "# 配置 PUBLIC_URL、MYSQL_*、STORAGE_BACKEND=minio、MINIO_*\n"
        "chmod +x start-app.sh\n"
        "./start-app.sh\n"
        "# 外挂 MySQL：可复制 docker-compose.app.external-db.example.yml\n"
        "# 为 docker-compose.app.override.yml，或设置 USE_EXTERNAL_MYSQL=1",
    )
    add_heading(doc, "6.3 机器 B（MinIO 初始化）", 2)
    add_code(
        doc,
        "cp .env.storage.example .env\n"
        "chmod +x start-storage.sh\n"
        "./start-storage.sh",
    )
    add_heading(doc, "6.4 常用运维命令", 2)
    add_code(
        doc,
        "# 查看状态 / 日志（机器 A）\n"
        "docker compose -f docker-compose.app.yml ps\n"
        "docker compose -f docker-compose.app.yml logs -f backend\n"
        "docker compose -f docker-compose.app.yml logs -f scheduler\n\n"
        "# 健康检查\n"
        "curl http://127.0.0.1/api/health/\n\n"
        "# 重建并启动\n"
        "./start-app.sh\n"
        "# 或\n"
        "docker compose -f docker-compose.app.yml up -d --build",
    )
    add_heading(doc, "6.5 备份建议", 2)
    add_bullets(
        doc,
        [
            "数据库：mysqldump（或 scripts/backup_mysql.py / backup_all.py）",
            "图片（MinIO）：mc mirror 备份 bucket 前缀 data/image_db/upload/",
            "图片（本地）：备份项目 upload/ 目录",
            "配置：妥善保管根目录 .env（含密钥）",
        ],
    )

    # 7
    add_heading(doc, "7. 本地开发环境", 1)
    add_heading(doc, "7.1 后端", 2)
    add_code(
        doc,
        "cd backend\n"
        "pip install -r requirements.txt\n"
        "copy .env.example .env   # Windows\n"
        "# DB_ENGINE=sqlite 可免 MySQL；或配置 mysql\n"
        "python manage.py migrate\n"
        "python manage.py runserver 0.0.0.0:8000\n"
        "# 健康检查 http://127.0.0.1:8000/api/health/\n"
        "# Swagger   http://127.0.0.1:8000/api/docs/",
    )
    add_heading(doc, "7.2 前端", 2)
    add_code(
        doc,
        "cd frontend\n"
        "npm install\n"
        "npm run dev\n"
        "# http://localhost:5173  ，/api 代理到 127.0.0.1:8000（vite.config.js）",
    )
    add_heading(doc, "7.3 生产前端构建", 2)
    add_code(
        doc,
        "cd frontend\n"
        "copy .env.production.example .env.production\n"
        "npm run build\n"
        "# 产物：frontend/dist/ ；Docker 构建时由 web 镜像完成",
    )
    add_para(
        doc,
        "Windows 注意：项目路径必须是纯 ASCII（如 E:\\image_db），含中文路径会导致 Docker BuildKit 失败。",
        bold=True,
    )

    # 8
    add_heading(doc, "8. 功能模块与关键文件", 1)
    add_heading(doc, "8.1 API 前缀", 2)
    add_table(
        doc,
        ["前缀", "模块"],
        [
            ["/api/auth/", "登录、刷新、当前用户"],
            ["/api/images/", "上传、浏览、BLOB 迁移/导出/同步、外库连接"],
            ["/api/fingerprints/", "指纹对、导入任务、业务评测"],
            ["/api/sql/", "管理员 SQL / 模板"],
            ["/api/logs/", "操作日志"],
            ["/api/config/", "系统设置"],
            ["/api/health/", "就绪检查（DB、存储可写、密钥等）"],
        ],
    )
    add_heading(doc, "8.2 前端页面 ↔ 路由", 2)
    add_table(
        doc,
        ["菜单/页面", "路由", "主要文件"],
        [
            ["首页", "/", "frontend/src/views/Home.vue"],
            ["图片上传", "/upload", "Upload.vue"],
            ["任务台", "/blob-migrate", "BlobMigrate.vue"],
            ["模拟数据库", "/blob-browse", "BlobTableViews.vue"],
            ["指纹对比", "/fingerprint-pairs", "FingerprintPairs.vue"],
            ["指纹评测", "/fingerprint-eval", "FingerprintEval.vue（菜单隐藏）"],
            ["分类", "/categories", "CategoryManage.vue"],
            ["操作日志", "/logs", "Logs.vue"],
            ["系统设置", "/settings", "Settings.vue"],
        ],
    )
    add_heading(doc, "8.3 后端关键实现", 2)
    add_table(
        doc,
        ["能力", "主要位置"],
        [
            ["BLOB 迁移任务", "backend/images/blob_migration_*"],
            ["模拟库/表视图配置", "backend/images/blob_table_view_*"],
            ["目录树/外库", "backend/images/blob_catalog_* 、 external_db_*"],
            ["路径导出任务", "export job 相关 views/commands"],
            ["指纹导入/对比/评测", "backend/fingerprints/"],
            ["模拟 SQL", "backend/sqlquery/simulated_sql.py"],
            ["存储抽象", "backend/utils/storage.py"],
            ["路径规范", "backend/utils/path_builder.py"],
        ],
    )
    add_heading(doc, "8.4 前端状态与轮询", 2)
    add_bullets(
        doc,
        [
            "路径导出任务：frontend/src/stores/backgroundExport.js",
            "指纹导入任务：frontend/src/stores/fingerprintImport.js",
            "页面数据刷新：frontend/src/utils/usePageDataRefresh.js",
            "隐藏标签页时会暂停轮询（visibilityAwarePoll），减轻主线程压力",
            "任务台集中展示进度；布局挂载时会延后 restoreFromSession",
        ],
    )

    # 9 — How to modify
    add_heading(doc, "9. 如何修改：常见开发场景", 1)
    add_para(doc, "以下是交接后最常被问到的「我该改哪里」。")

    add_heading(doc, "9.1 新增一个后端 API", 2)
    add_numbered(
        doc,
        [
            "在对应 app 下新增或扩展视图（如 backend/images/views.py 或独立 *_views.py）",
            "在该 app 的 urls.py 注册路由（images/urls.py、fingerprints/urls.py 等）",
            "确认已挂到 backend/config/urls.py（一般已 include）",
            "在 frontend/src/api/ 对应模块增加封装函数",
            "页面中调用；必要时补 Django 单测",
        ],
    )
    add_code(
        doc,
        "cd backend\n"
        "set DB_ENGINE=sqlite          # PowerShell: $env:DB_ENGINE=\"sqlite\"\n"
        "python manage.py test images fingerprints sqlquery -v 1 --noinput",
    )

    add_heading(doc, "9.2 新增一个前端页面并挂菜单", 2)
    add_numbered(
        doc,
        [
            "新建 frontend/src/views/MyPage.vue",
            "在 frontend/src/config/menu.js 的 MENU_ITEMS 增加一项（path/name/title/icon/adminOnly）",
            "在 frontend/src/router/index.js 的 VIEW_MAP / 路由表注册组件",
            "若使用新图标，在 MainLayout.vue 的 iconMap 中补充 Element Plus 图标",
            "本地 npm run dev 验证；Docker 部署需重建 web 镜像后浏览器强刷（Ctrl+F5）",
        ],
    )

    add_heading(doc, "9.3 修改侧边栏文案/顺序/权限", 2)
    add_para(doc, "只需改 frontend/src/config/menu.js。adminOnly: true 的项仅管理员可见；路由守卫在 frontend/src/router/guards.js。")

    add_heading(doc, "9.4 数据库表结构变更", 2)
    add_bullets(
        doc,
        [
            "Django 自管表：python manage.py makemigrations && migrate",
            "大量业务表使用 managed=False，结构升级常走启动时 schema_ensure（见 fingerprints/schema_ensure.py、entrypoint）",
            "参考 DDL 在 sql/ 目录；生产变更前先备份",
        ],
    )

    add_heading(doc, "9.5 改 Docker / Nginx / 启动逻辑", 2)
    add_table(
        doc,
        ["想改什么", "改哪里"],
        [
            ["服务组成、端口映射", "docker-compose.yml 或 docker-compose.app.yml"],
            ["后端镜像与依赖", "docker/Dockerfile.backend、backend/requirements*.txt"],
            ["前端构建与 Nginx", "docker/Dockerfile.nginx、docker/nginx/default.conf"],
            ["容器启动：迁移/回收任务", "docker/backend-entrypoint.sh"],
            ["定时轮询与维护", "docker/maintenance-loop.sh"],
            ["PUBLIC_URL→Hosts/CORS", "docker/set-env.py"],
        ],
    )

    add_heading(doc, "9.6 切换存储：本地 ↔ MinIO", 2)
    add_numbered(
        doc,
        [
            "在 .env 设置 STORAGE_BACKEND=local 或 minio，并填好 MINIO_*",
            "实现位于 backend/utils/storage.py",
            "容器内验证：docker compose exec backend python scripts/init_storage.py",
            "健康检查关注 readiness.upload_writable",
        ],
    )

    add_heading(doc, "9.7 改上传限制 / SQL 超时等系统参数", 2)
    add_para(
        doc,
        "管理员登录 → 系统设置页面；或调用 /api/config/。"
        "部分参数也可在 settings.py / .env 中有默认值，以页面与数据库配置为准时注意覆盖关系。",
    )

    add_heading(doc, "9.8 改图片路径规则", 2)
    add_para(
        doc,
        "路径约定见 docs/storage.md；代码入口 backend/utils/path_builder.py。"
        "相对路径形态一般为：upload/{YYYYMMDD}/{category_id}/{uuid}.ext；"
        "MinIO 对象键会再加 MINIO_PREFIX。",
    )

    add_heading(doc, "9.9 指纹相关改动入口", 2)
    add_bullets(
        doc,
        [
            "说明文档：docs/fingerprint_pairs.md",
            "后端：backend/fingerprints/",
            "前端对比页：FingerprintPairs.vue；评测：FingerprintEval.vue",
            "导入任务回收：manage.py reclaim_fingerprint_import_jobs（scheduler 会 kick）",
        ],
    )

    add_heading(doc, "9.10 前端交互与错误提示约定", 2)
    add_bullets(
        doc,
        [
            "全局请求封装：frontend/src/api/request.js（JWT、统一错误）",
            "业务错误 toast：优先 showRequestError，避免重复弹窗（__globalToastShown）",
            "可重试请求：callWithRetry（重试中途会抑制全局错误 toast）",
        ],
    )

    # 10
    add_heading(doc, "10. 后台任务、调度与故障处理", 1)
    add_heading(doc, "10.1 scheduler 在做什么", 2)
    add_para(doc, "scheduler 容器循环执行 docker/maintenance-loop.sh：")
    add_bullets(
        doc,
        [
            "周期性：处理 BLOB 迁移队列、路径导出、外库同步、指纹导入 kick",
            "按 MAINTENANCE_INTERVAL_HOURS：清理过期日志与软删文件",
            "backend 容器内还有导出 sidecar 循环，避免长导出阻塞 Gunicorn 启动",
        ],
    )
    add_heading(doc, "10.2 重启后任务卡在 running", 2)
    add_para(doc, "entrypoint / scheduler 会执行 reclaim_* 命令，把卡住的 running 收回 pending。也可手动：")
    add_code(
        doc,
        "docker compose -f docker-compose.app.yml exec backend \\\n"
        "  python manage.py reclaim_blob_migration_jobs\n"
        "docker compose -f docker-compose.app.yml exec backend \\\n"
        "  python manage.py reclaim_blob_export_jobs --include-paused\n"
        "docker compose -f docker-compose.app.yml exec backend \\\n"
        "  python manage.py reclaim_fingerprint_import_jobs",
    )
    add_heading(doc, "10.3 常见故障速查", 2)
    add_table(
        doc,
        ["现象", "排查/处理"],
        [
            ["启动后 502", "等 MySQL 初始化约 1 分钟；看 backend 日志"],
            ["登录 403", "PUBLIC_URL 与浏览器地址不一致；改 .env 后重启"],
            ["模拟库「已加载 0」", "可能连到空的 compose db；跑 scripts/diagnose-machine-a.sh，确认 DB_HOST"],
            ["端口占用", ".env 设 HTTP_PORT=8080"],
            ["前端仍是旧界面", "重建 web 镜像 + 浏览器 Ctrl+F5"],
            ["MinIO 不可写", "curl VIP:9000/minio/health/live；exec init_storage.py"],
            ["标题栏最小化不灵敏", "多为页面主线程忙碌；无任务时也应减少空轮询（已做可见性暂停）"],
        ],
    )

    # 11
    add_heading(doc, "11. 测试与验收", 1)
    add_heading(doc, "11.1 单元测试", 2)
    add_code(
        doc,
        "cd backend\n"
        "$env:DB_ENGINE=\"sqlite\"   # PowerShell\n"
        "python manage.py test images sqlquery fingerprints -v 1 --noinput",
    )
    add_heading(doc, "11.2 冒烟测试", 2)
    add_code(
        doc,
        "python scripts/smoke_test.py --base-url http://127.0.0.1\n"
        "python scripts/smoke_test.py --base-url http://<服务器IP> --username admin --password <密码>",
    )
    add_para(doc, "仓库当前未配置 .github/workflows CI，发布前请人工跑测试与冒烟。")
    add_heading(doc, "11.3 建议验收路径", 2)
    add_numbered(
        doc,
        [
            "/api/health/ 返回就绪",
            "admin 登录成功，改密",
            "上传一张图并可预览",
            "模拟数据库：连外库 → 建配置 → 浏览行 → SQL 查询",
            "任务台：发起一小批迁移或导出，观察进度与完成",
            "指纹对比：导入或浏览样本，打开叠加对比",
            "管理员查看操作日志与系统设置",
        ],
    )

    # 12
    add_heading(doc, "12. 安全与注意事项", 1)
    add_bullets(
        doc,
        [
            "切勿将 .env、真实密钥、生产上传目录提交到 Git",
            "默认 admin/admin123 仅用于初始化",
            "图片不直接静态暴露，统一鉴权访问（见 docs/file_security.md）",
            "管理员 SQL 仅允许受控查询；注意 SQL 超时配置",
            "生产建议关闭 Django DEBUG，使用强 SECRET_KEY / IMAGE_ACCESS_SECRET",
            "外库连接密码保存在系统库中，交接时说明权限边界与可连网段",
        ],
    )

    # 13
    add_heading(doc, "13. 交接检查清单", 1)
    add_table(
        doc,
        ["项", "状态（交接时勾选）", "备注"],
        [
            ["Git 仓库权限（读/写）", "□", "github.com/luoshirui-192/image"],
            ["机器 A 服务器登录方式", "□", "SSH / 堡垒机"],
            ["机器 B / MinIO 凭据与 VIP", "□", "192.168.9.9:9000"],
            ["根目录 .env 已移交（不入库）", "□", "含 SECRET / DB / MinIO"],
            ["PUBLIC_URL 与实际访问一致", "□", ""],
            ["默认管理员密码已修改", "□", ""],
            ["备份策略已说明", "□", "DB + 对象存储"],
            ["外挂 MySQL（若有）主机与库名", "□", "如 mysql8039"],
            ["能独立完成 start-app.sh 重启", "□", ""],
            ["能完成一次冒烟测试", "□", "smoke_test.py"],
            ["知道如何加菜单/加 API", "□", "见第 9 章"],
            ["知道 reclaim 卡住任务", "□", "见第 10 章"],
        ],
    )

    add_heading(doc, "附录 A：相关文档索引", 1)
    add_table(
        doc,
        ["文档", "内容"],
        [
            ["README.md", "总览与单机 Docker"],
            ["README-MACHINE-A.md", "机器 A"],
            ["README-MACHINE-B.md", "机器 B / MinIO"],
            ["docs/deploy.md", "裸机 Nginx + Gunicorn"],
            ["docs/quickstart-docker.md", "Docker 速查"],
            ["docs/storage.md", "路径与存储约定"],
            ["docs/fingerprint_pairs.md", "指纹功能"],
            ["docs/file_security.md", "文件安全"],
            ["docs/github.md", "仓库协作说明"],
            ["backend/README.md / frontend/README.md", "前后端补充说明"],
        ],
    )

    add_heading(doc, "附录 B：快速命令备忘", 1)
    add_code(
        doc,
        "# 机器 A 重启\n./start-app.sh\n\n"
        "# 日志\ndocker compose -f docker-compose.app.yml logs -f backend scheduler\n\n"
        "# 健康\ncurl http://127.0.0.1/api/health/\n\n"
        "# 本地开发\ncd backend && python manage.py runserver 0.0.0.0:8000\n"
        "cd frontend && npm run dev\n\n"
        "# 测试\ncd backend && set DB_ENGINE=sqlite && python manage.py test images -v 1\n"
        "python scripts/smoke_test.py --base-url http://127.0.0.1",
    )

    add_para(doc, "— 文档结束。后续功能变更请同步更新本交接文档或 README。", bold=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
